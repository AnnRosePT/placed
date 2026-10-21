"""Analytics & Reports endpoints (Phase 10)."""
from __future__ import annotations

import io
import math
from datetime import date
from typing import Any

import numpy as np
import pandas as pd
from flask import Blueprint, Response, g, jsonify, request

from ..auth.decorators import require_auth
from ..services.audit import log
from ..services.data import (
    DEPARTMENTS, SKILL_AXES, FilterError, actor_id, fetch_students, latest_by_student, parse_filters,
)
from ..services.settings_service import get_settings, num
from ..utils.csv_utils import csv_response, safe_cell
from ..utils.errors import error_response

analytics_bp = Blueprint("analytics", __name__, url_prefix="/api")

CGPA_TIERS = [("> 9.0", 9.0, 10.01), ("8.0 – 8.9", 8.0, 9.0), ("7.0 – 7.9", 7.0, 8.0), ("< 7.0", 0.0, 7.0)]


def _filters():
    return parse_filters(request.args)


def _frame(filters) -> pd.DataFrame:
    students = fetch_students(filters)
    cols = ["id", "usn", "full_name", "department", "batch_year", "cgpa", "placement_status", "company",
            "package_lpa", "aptitude_score", "technical_score", "communication_score"]
    frame = pd.DataFrame(students, columns=cols) if students else pd.DataFrame(columns=cols)
    for c in ("cgpa", "package_lpa"):
        frame[c] = pd.to_numeric(frame[c], errors="coerce")
    return frame


def _round(value, digits=2):
    if value is None or (isinstance(value, float) and (math.isnan(value) or math.isinf(value))):
        return None
    return round(float(value), digits)


def _with_filters(view):
    from functools import wraps

    @wraps(view)
    def wrapped(*args, **kwargs):
        try:
            filters = _filters()
        except FilterError as exc:
            return error_response(400, "validation_error", str(exc))
        return view(filters, *args, **kwargs)
    return wrapped


def _dept_rows(frame: pd.DataFrame, dream: float) -> list[dict[str, Any]]:
    rows = []
    for dept in DEPARTMENTS:
        sub = frame[frame["department"] == dept]
        if not len(sub):
            continue
        eligible = sub[sub["placement_status"] != "opted_out"]
        placed = sub[sub["placement_status"] == "placed"]
        pk = placed["package_lpa"].dropna()
        rows.append({
            "department": dept, "enrolled": int(len(sub)), "eligible": int(len(eligible)),
            "placed": int(len(placed)), "rate": _round(len(placed) / len(eligible), 4) if len(eligible) else None,
            "avg_package_lpa": _round(pk.mean()) if len(pk) else None,
            "top_package_lpa": _round(pk.max()) if len(pk) else None,
            "dream_offers": int((pk >= dream).sum()),
        })
    return rows


@analytics_bp.get("/analytics/overview")
@require_auth
@_with_filters
def overview(filters):
    settings = get_settings()
    dream = num(settings, "dream_lpa")
    frame = _frame(filters)
    eligible = frame[frame["placement_status"] != "opted_out"]
    placed = frame[frame["placement_status"] == "placed"]
    pk = placed["package_lpa"].dropna()
    depts = [d for d in _dept_rows(frame, dream) if d["rate"] is not None]
    rates = [d["rate"] for d in depts]
    top = placed.loc[placed["package_lpa"].idxmax()] if len(pk) else None
    return jsonify({
        "filters": filters,
        "students": int(len(frame)),
        "eligible": int(len(eligible)),
        "placed": int(len(placed)),
        "conversion_rate": _round(len(placed) / len(eligible), 4) if len(eligible) else None,
        "target_rate": num(settings, "placement_target_rate"),
        "median_package_lpa": _round(pk.median()) if len(pk) else None,
        "mean_package_lpa": _round(pk.mean()) if len(pk) else None,
        "max_package_lpa": _round(pk.max()) if len(pk) else None,
        "max_package_company": None if top is None else top["company"],
        "inter_branch_parity": _round(min(rates) / max(rates), 4) if rates and max(rates) > 0 else None,
        "parity_low": min(depts, key=lambda d: d["rate"])["department"] if depts else None,
        "parity_high": max(depts, key=lambda d: d["rate"])["department"] if depts else None,
        "tier1_share": _round((pk >= dream).sum() / len(placed), 4) if len(placed) else None,
        "dream_lpa": dream,
        "batches": sorted({int(s["batch_year"]) for s in fetch_students({}) if s.get("batch_year")}),
    })


@analytics_bp.get("/analytics/departments")
@require_auth
@_with_filters
def departments(filters):
    dream = num(get_settings(), "dream_lpa")
    return jsonify({"items": _dept_rows(_frame(filters), dream), "dream_lpa": dream})


def point_biserial(frame: pd.DataFrame) -> tuple[float | None, int]:
    """Pearson r between CGPA and a 0/1 placed outcome (= point-biserial correlation).

    Uses students with a final outcome (placed / not_placed) and a CGPA.
    """
    decided = frame[frame["placement_status"].isin(["placed", "not_placed"])].dropna(subset=["cgpa"])
    if len(decided) < 3:
        return None, int(len(decided))
    y = (decided["placement_status"] == "placed").astype(float).to_numpy()
    x = decided["cgpa"].astype(float).to_numpy()
    if y.std() == 0 or x.std() == 0:
        return None, int(len(decided))
    return float(np.corrcoef(x, y)[0, 1]), int(len(decided))


@analytics_bp.get("/analytics/cgpa-correlation")
@require_auth
@_with_filters
def cgpa_correlation(filters):
    frame = _frame(filters)
    tiers = []
    for label, low, high in CGPA_TIERS:
        sub = frame[(frame["cgpa"] >= low) & (frame["cgpa"] < high)]
        eligible = sub[sub["placement_status"] != "opted_out"]
        placed = sub[sub["placement_status"] == "placed"]
        pk = placed["package_lpa"].dropna()
        tiers.append({"label": label, "students": int(len(sub)), "placed": int(len(placed)),
                      "rate": _round(len(placed) / len(eligible), 4) if len(eligible) else None,
                      "avg_package_lpa": _round(pk.mean()) if len(pk) else None})
    r, n = point_biserial(frame)
    latest = latest_by_student("predictions", list(frame["id"]), "student_id, probability, created_at") if len(frame) else {}
    points = [{"id": row["id"], "usn": row["usn"], "full_name": row["full_name"], "cgpa": _round(row["cgpa"]),
               "probability": float(latest[row["id"]]["probability"]), "status": row["placement_status"]}
              for _, row in frame.iterrows() if row["id"] in latest and not pd.isna(row["cgpa"])]
    return jsonify({"tiers": tiers, "pearson_r": _round(r, 4), "n": n, "points": points})


def _package_tier(value: float, s: dict[str, Any]) -> str:
    if value >= num(s, "super_dream_lpa"):
        return "super_dream"
    if value >= num(s, "dream_lpa"):
        return "dream"
    if value >= num(s, "high_potential_lpa"):
        return "high_potential"
    return "core_it"


@analytics_bp.get("/analytics/packages")
@require_auth
@_with_filters
def packages(filters):
    settings = get_settings()
    frame = _frame(filters)
    pk = frame[frame["placement_status"] == "placed"]["package_lpa"].dropna()
    bins = []
    if len(pk):
        width = 2.0
        top = math.ceil(pk.max() / width) * width or width
        edges = np.arange(0, top + width, width)
        counts, _ = np.histogram(pk, bins=edges)
        bins = [{"from": float(edges[i]), "to": float(edges[i + 1]), "count": int(c)} for i, c in enumerate(counts)]
    tiers = {"super_dream": 0, "dream": 0, "high_potential": 0, "core_it": 0}
    for v in pk:
        tiers[_package_tier(float(v), settings)] += 1
    return jsonify({
        "values": [float(v) for v in pk], "bins": bins,
        "tiers": [{"key": k, "count": v, "share": _round(v / len(pk), 4) if len(pk) else None} for k, v in tiers.items()],
        "cutoffs": {k: num(settings, k) for k in ("super_dream_lpa", "dream_lpa", "high_potential_lpa")},
    })


@analytics_bp.get("/analytics/skills")
@require_auth
@_with_filters
def skills(filters):
    students = fetch_students(filters)
    latest = latest_by_student("skill_assessments", [s["id"] for s in students], "*", "assessed_at") if students else {}
    dept_of = {s["id"]: s["department"] for s in students}
    rows = [{"department": dept_of[sid], **{a: a_row.get(a) for a in SKILL_AXES}} for sid, a_row in latest.items()]
    items = []
    if rows:
        frame = pd.DataFrame(rows)
        for dept, sub in frame.groupby("department"):
            items.append({"department": dept, "assessed": int(len(sub)),
                          **{a: _round(pd.to_numeric(sub[a], errors="coerce").mean(), 1) for a in SKILL_AXES}})
    items.sort(key=lambda d: DEPARTMENTS.index(d["department"]) if d["department"] in DEPARTMENTS else 99)
    return jsonify({"axes": list(SKILL_AXES), "items": items})


# ── Report exports ───────────────────────────────────────────────────────
def _record_export(title: str, kind: str, scope: dict, row_count: int) -> None:
    g.db.table("report_exports").insert({"title": title, "type": kind, "scope": scope,
                                         "row_count": row_count, "created_by": actor_id()}).execute()
    log("export", "report_exports", None, {"title": title, "type": kind, "row_count": row_count, "scope": scope})


@analytics_bp.get("/reports/department-ledger.csv")
@require_auth
@_with_filters
def department_ledger(filters):
    dream = num(get_settings(), "dream_lpa")
    rows = _dept_rows(_frame(filters), dream)
    _record_export("Department placement ledger", "csv", filters, len(rows))
    header = ["department", "enrolled", "eligible", "placed", "rate", "avg_package_lpa", "top_package_lpa",
              f"dream_offers_ge_{dream:g}_lpa"]
    return csv_response(header, ([r["department"], r["enrolled"], r["eligible"], r["placed"], r["rate"],
                                  r["avg_package_lpa"], r["top_package_lpa"], r["dream_offers"]] for r in rows),
                        "department_ledger")


@analytics_bp.get("/reports/company-audit.xlsx")
@require_auth
@_with_filters
def company_audit(filters):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    frame = _frame(filters)
    placed = frame[(frame["placement_status"] == "placed") & frame["company"].notna() & (frame["company"] != "")]
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    header_font = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1E1B4B")
    summary.append(["Company", "Hires", "Avg package (LPA)", "Max package (LPA)"])
    for cell in summary[1]:
        cell.font, cell.fill = header_font, fill
    used_titles = {"Summary"}
    for company, sub in sorted(placed.groupby("company"), key=lambda item: -len(item[1])):
        pk = sub["package_lpa"].dropna()
        summary.append([safe_cell(company), int(len(sub)), _round(pk.mean()) if len(pk) else None,
                        _round(pk.max()) if len(pk) else None])
        title = "".join(ch for ch in str(company) if ch not in '[]:*?/\\')[:28] or "Company"
        base, n = title, 2
        while title in used_titles:
            title = f"{base[:25]}-{n}"
            n += 1
        used_titles.add(title)
        sheet = workbook.create_sheet(title)
        sheet.append(["USN", "Name", "Department", "Batch", "CGPA", "Package (LPA)"])
        for cell in sheet[1]:
            cell.font, cell.fill = header_font, fill
        for _, row in sub.iterrows():
            sheet.append([safe_cell(row["usn"]), safe_cell(row["full_name"]), row["department"],
                          None if pd.isna(row["batch_year"]) else int(row["batch_year"]),
                          _round(row["cgpa"]), _round(row["package_lpa"])])
        for col in "ABCDEF":
            sheet.column_dimensions[col].width = 18
    for col in "ABCD":
        summary.column_dimensions[col].width = 22
    buffer = io.BytesIO()
    workbook.save(buffer)
    _record_export("Company recruitment audit", "xlsx", filters, int(len(placed)))
    return Response(buffer.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             f'attachment; filename="company_audit_{date.today().isoformat()}.xlsx"'})


@analytics_bp.get("/reports/students.csv")
@require_auth
@_with_filters
def student_outcomes(filters):
    frame = _frame(filters)
    _record_export("Student outcome register", "csv", filters, int(len(frame)))
    cols = ["usn", "full_name", "department", "batch_year", "cgpa", "aptitude_score", "technical_score",
            "communication_score", "placement_status", "company", "package_lpa"]
    return csv_response(cols, (["" if pd.isna(v) else v for v in row] for row in frame[cols].itertuples(index=False)),
                        "student_outcomes")


@analytics_bp.get("/reports/history")
@require_auth
def report_history():
    rows = g.db.table("report_exports").select("*").order("created_at", desc=True).limit(25).execute().data or []
    return jsonify({"items": rows})
